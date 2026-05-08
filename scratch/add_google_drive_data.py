import openpyxl
import os

file_path = r'testdata\excel\recruitment_data.xlsx'

def add_test_data():
    if not os.path.exists(file_path):
        print(f"Error: {file_path} not found")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    # Get current headers
    headers = [cell.value for cell in sheet[1]]
    print(f"Current headers: {headers}")

    # Required headers for the new scenario
    required_headers = ['source', 'google_drive_link', 'candidate_name']
    
    # Add missing headers
    for h in required_headers:
        if h not in headers:
            headers.append(h)
            sheet.cell(row=1, column=len(headers)).value = h
            print(f"Added header: {h}")

    # Prepare data row
    data = {
        'scenario': 'Connect Google Drive and Verify Candidate Profile',
        'role': 'Admin', # Default role
        'job_role': 'QA Engineer',
        'source': 'Google Drive',
        'google_drive_link': 'https://drive.google.com/drive/folders/1McKz4CeA8k6TcQnjtseBzJ15Sk1k0hSE?usp=sharing',
        'candidate_name': 'Anand Tester' # Example name
    }

    # Find the next empty row
    next_row = sheet.max_row + 1
    
    # Fill the row based on headers
    for i, header in enumerate(headers, 1):
        if header in data:
            sheet.cell(row=next_row, column=i).value = data[header]

    try:
        wb.save(file_path)
        print(f"Successfully added test data to row {next_row} in {file_path}")
    except PermissionError:
        new_path = file_path.replace('.xlsx', '_updated.xlsx')
        wb.save(new_path)
        print(f"Permission denied for {file_path}. Saved data to {new_path} instead. Please close Excel and merge if necessary.")

if __name__ == "__main__":
    add_test_data()
