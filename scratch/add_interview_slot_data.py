import openpyxl
import os

file_path = r'testdata\excel\recruitment_data.xlsx'

def add_interview_slot_data():
    if not os.path.exists(file_path):
        print(f"Error: {file_path} not found")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    # Get current headers
    headers = [cell.value for cell in sheet[1]]
    print(f"Current headers: {headers}")

    # Required headers for the new scenario
    required_headers = ['start_date', 'end_date', 'interviewer_name']
    
    # Add missing headers
    for h in required_headers:
        if h not in headers:
            headers.append(h)
            sheet.cell(row=1, column=len(headers)).value = h
            print(f"Added header: {h}")

    # Prepare data row
    data = {
        'scenario': 'Add interview slot manually and verify date and time',
        'role': 'Admin',
        'job_role': 'QA Engineer',
        'start_date': 'May 8, 2026 08:00 AM',
        'end_date': 'May 8, 2026 09:00 AM',
        'interviewer_name': 'John Doe'
    }

    # Find the next empty row
    next_row = sheet.max_row + 1
    
    # Fill the row based on headers
    for i, header in enumerate(headers, 1):
        if header in data:
            sheet.cell(row=next_row, column=i).value = data[header]

    import time
    for attempt in range(5):
        try:
            wb.save(file_path)
            print(f"Successfully added test data to row {next_row} in {file_path}")
            return
        except PermissionError:
            print(f"Attempt {attempt+1}: Permission denied for {file_path}. Retrying in 2s...")
            time.sleep(2)
    print(f"Failed to save {file_path} after 5 attempts. Please close Excel.")

if __name__ == "__main__":
    add_interview_slot_data()
