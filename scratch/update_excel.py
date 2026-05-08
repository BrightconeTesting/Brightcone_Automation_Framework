import openpyxl
import os

def update_candidate_to_available():
    file_path = r'c:\Users\Dell\OneDrive\Desktop\Bright_Automation_Before_Adding_Session\testdata\excel\recruitment_data.xlsx'
    
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    headers = [cell.value for cell in sheet[1]]
    
    scenario_to_update = "send interview invitation"
    candidate_name_new = "PINJARI SHAMEENA"
    
    scenario_col = headers.index("scenario") + 1
    candidate_name_col = headers.index("candidate_name") + 1
    
    updated = False
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=scenario_col).value == scenario_to_update:
            sheet.cell(row=row, column=candidate_name_col).value = candidate_name_new
            updated = True
            print(f"Updated row {row} with candidate name '{candidate_name_new}'.")
    
    if not updated:
        print("Scenario not found in Excel.")
    else:
        wb.save(file_path)
        print("Excel updated successfully.")

if __name__ == "__main__":
    update_candidate_to_available()
