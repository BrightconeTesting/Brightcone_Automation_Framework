import openpyxl
import os

file_path = r'testdata\excel\recruitment_data.xlsx'
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

headers = [cell.value for cell in sheet[1]]

new_row_data = {
    'scenario': 'Submit feedback and verify',
    'role': 'admin',
    'expected_result': 'HIRE',
    'job_role': 'QA Engineer'
}

next_row = sheet.max_row + 1

for i, header in enumerate(headers, 1):
    if header in new_row_data:
        sheet.cell(row=next_row, column=i).value = new_row_data[header]

wb.save(file_path)
print("Feedback test data added successfully.")
