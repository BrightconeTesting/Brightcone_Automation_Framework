import openpyxl
import os

file_path = r'testdata\excel\recruitment_data.xlsx'
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

# Get headers to ensure correct column placement
headers = [cell.value for cell in sheet[1]]

new_row_data = {
    'scenario': 'Cancel scheduled interview and verify status',
    'role': 'admin',
    'expected_result': 'CANCELLED',
    'job_role': 'QA Engineer'
}

# Find the next empty row
next_row = sheet.max_row + 1

# Fill the row based on headers
for i, header in enumerate(headers, 1):
    if header in new_row_data:
        sheet.cell(row=next_row, column=i).value = new_row_data[header]

wb.save(file_path)
print("Cancellation test data added successfully.")
