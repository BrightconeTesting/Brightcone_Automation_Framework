import openpyxl
import os

def update_date_formats():
    file_path = r'testdata\excel\recruitment_data.xlsx'
    
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    headers = [cell.value for cell in sheet[1]]
    
    scenario_col = headers.index("scenario") + 1
    start_date_col = headers.index("start_date") + 1
    end_date_col = headers.index("end_date") + 1
    
    updates = {
        'Add interview slot manually and verify date and time': {
            'start_date': 'May 8, 2026 08:00 AM',
            'end_date': 'May 8, 2026 09:00 AM'
        },
        'Re-schedule interview slot and verify status': {
            'start_date': 'May 8, 2026 08:00 AM',
            'end_date': 'May 8, 2026 09:00 AM'
        },
        'Sync interview slot from calendar and verify date and time': {
            'from_date': 'May 8, 2026 08:00 AM',
            'to_date': 'May 8, 2026 09:00 AM'
        }
    }
    
    updated_count = 0
    for row in range(2, sheet.max_row + 1):
        scenario_name = sheet.cell(row=row, column=scenario_col).value
        if scenario_name in updates:
            scenario_updates = updates[scenario_name]
            for key, value in scenario_updates.items():
                if key in headers:
                    col_idx = headers.index(key) + 1
                    sheet.cell(row=row, column=col_idx).value = value
            updated_count += 1
            print(f"Updated scenario: {scenario_name}")
    
    if updated_count == 0:
        print("No matching scenarios found to update.")
    else:
        wb.save(file_path)
        print(f"Excel updated successfully with {updated_count} updates.")

if __name__ == "__main__":
    update_date_formats()
